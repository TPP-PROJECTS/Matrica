from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from openpyxl import load_workbook


SOURCE_DOCX = Path("/Volumes/T9/Research/Research-AI/paper/Draft V3.docx")
SOURCE_XLSX = Path(
    "/Volumes/T9/Research/Research-AI/resource/"
    "Matrica_Competitor_Matrix_All_2026-08-26.xlsx"
)
ASSET_DIR = Path("/private/tmp/draftv3-interface-assets.Xa177g")
OUTPUT_DOCX = Path(
    "/Volumes/T9/TPP/topic-generator/docs/"
    "Draft_V3_Current_Technical_Appendices_Highlighted_2026-09-01.docx"
)

YELLOW = WD_COLOR_INDEX.YELLOW


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=55, start=55, bottom=55, end=55) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_cell_width(cell, width_inches: float) -> None:
    cell.width = Inches(width_inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def highlight_run(run, *, bold: bool | None = None, italic: bool | None = None, size=None):
    run.font.highlight_color = YELLOW
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if size is not None:
        run.font.size = Pt(size)
    return run


def add_highlighted_paragraph(
    doc: Document,
    text: str = "",
    *,
    style: str | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
    alignment=None,
    keep_with_next: bool = False,
    size: float | None = None,
):
    paragraph = doc.add_paragraph(style=style)
    if alignment is not None:
        paragraph.alignment = alignment
    paragraph.paragraph_format.keep_with_next = keep_with_next
    if text:
        highlight_run(paragraph.add_run(text), bold=bold, italic=italic, size=size)
    return paragraph


def clear_cell(cell) -> None:
    paragraph = cell.paragraphs[0]
    for child in list(paragraph._p):
        paragraph._p.remove(child)


def add_cell_text(
    cell,
    text: str,
    *,
    bold: bool = False,
    size: float = 8.0,
    align=WD_ALIGN_PARAGRAPH.LEFT,
) -> None:
    clear_cell(cell)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    highlight_run(paragraph.add_run(text), bold=bold, size=size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def add_hyperlink(cell, text: str, url: str, *, size: float = 8.0) -> None:
    clear_cell(cell)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    if not url or url in {"N/A", "—"}:
        highlight_run(paragraph.add_run(text), size=size)
        return

    relation_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)

    new_run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")

    run_style = OxmlElement("w:rStyle")
    run_style.set(qn("w:val"), "Hyperlink")
    run_properties.append(run_style)

    highlight = OxmlElement("w:highlight")
    highlight.set(qn("w:val"), "yellow")
    run_properties.append(highlight)

    font_size = OxmlElement("w:sz")
    font_size.set(qn("w:val"), str(int(size * 2)))
    run_properties.append(font_size)
    font_size_complex = OxmlElement("w:szCs")
    font_size_complex.set(qn("w:val"), str(int(size * 2)))
    run_properties.append(font_size_complex)

    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(run_properties)
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_heading(doc: Document, text: str, level: int):
    style = f"Heading {level}"
    paragraph = add_highlighted_paragraph(
        doc,
        text,
        style=style,
        bold=True,
        keep_with_next=True,
    )
    return paragraph


def add_figure(
    doc: Document,
    *,
    subsection: str,
    intro: str,
    image_path: Path,
    caption: str,
    final: bool = False,
) -> None:
    add_heading(doc, subsection, 2)
    add_highlighted_paragraph(doc, intro)
    picture_paragraph = doc.add_paragraph()
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_paragraph.paragraph_format.space_before = Pt(3)
    picture_paragraph.paragraph_format.space_after = Pt(3)
    picture_paragraph.add_run().add_picture(str(image_path), width=Inches(6.25))
    add_highlighted_paragraph(
        doc,
        caption,
        style="Caption",
        italic=True,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        keep_with_next=False,
        size=8.5,
    )
    if not final:
        doc.add_page_break()


def make_table(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    widths: list[float],
    *,
    font_size: float,
    hyperlink_column: int | None = None,
    hyperlinks: list[str] | None = None,
) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    prevent_row_split(header_row)
    for index, (cell, header, width) in enumerate(zip(header_row.cells, headers, widths)):
        set_table_cell_width(cell, width)
        set_cell_shading(cell, "D9EAF7")
        add_cell_text(
            cell,
            header,
            bold=True,
            size=font_size,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )

    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for column_index, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
            set_table_cell_width(cell, width)
            if row_index % 2:
                set_cell_shading(cell, "F5F8FA")
            if hyperlink_column is not None and column_index == hyperlink_column:
                add_hyperlink(
                    cell,
                    value,
                    hyperlinks[row_index] if hyperlinks else "",
                    size=font_size,
                )
            else:
                add_cell_text(cell, value, size=font_size)


def get_value(sheet, row: int, header_map: dict[str, int], name: str) -> str:
    value = sheet.cell(row, header_map[name]).value
    if value is None or str(value).strip() == "":
        return "—"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def add_appendix_a(doc: Document) -> None:
    doc.add_page_break()
    add_heading(doc, "Appendix A. Matrica Interface and Workflow", 1)
    add_highlighted_paragraph(
        doc,
        "This appendix links the workflow and implementation mechanisms described in "
        "Sections 5 and 6 to their visible interface controls. The screenshots are "
        "illustrative project views captured from the Matrica interface; they document "
        "how researchers move through the structured workflow and inspect or revise "
        "AI-assisted outputs. They do not extend the implementation or evaluation "
        "claims made in the main text.",
    )

    add_heading(doc, "A.1 Interface Map", 2)
    interface_rows = [
        ["1", "Chat & Define Topic", "Research Profile fields and topic-framing dialogue"],
        ["2", "Matrix Initial", "Researcher-controlled literature and system matrix schemas"],
        ["3", "System Review", "Selected comparison systems and researcher additions"],
        ["4", "Explore & Search", "Topics, Boolean queries, search limits, and database results"],
        ["5", "Select Paper", "Protocol, include/later/exclude decisions, reasons, and citation chasing"],
        ["6", "Matrix Table", "Paper/system rows, custom columns, AI-origin metadata, and evidence inspection"],
        ["7", "Synthesis", "Theme coverage, linked matrix dimensions, and editable analyses"],
        ["8", "Gap & Statement", "Coverage gaps, selected gap dimensions, and edited gap statement"],
        ["9", "Final Question", "Chosen topic, research question, study design, and paper structure"],
    ]
    make_table(
        doc,
        ["Stage", "Interface area", "Principal record or control"],
        interface_rows,
        [0.55, 1.45, 4.5],
        font_size=8.0,
    )
    add_highlighted_paragraph(
        doc,
        "Table A1. Interface stages and the principal research record or control exposed at each stage.",
        style="Caption",
        italic=True,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        size=8.5,
    )
    doc.add_page_break()

    figure_specs = [
        (
            "A.2.1 Structured Intent and Shared Workflow Shell",
            "The opening workspace captures the research profile before later stages are used. "
            "The persistent stage navigation keeps topic definition, matrices, search, screening, "
            "synthesis, gap analysis, and final question development within one project context.",
            "image4.png",
            "Figure A1. Research Profile and the persistent nine-stage workflow shell.",
        ),
        (
            "A.2.2 Protocol-Based Screening Workspace",
            "The screening workspace places protocol fields and include, later, or exclude actions "
            "beside each candidate record. Researcher decisions and reasons are retained as "
            "structured workflow records rather than disappearing into a chat transcript.",
            "image5.png",
            "Figure A2. Screening workspace with protocol controls and explicit researcher decisions.",
        ),
        (
            "A.2.3 Evidence Matrix and Inspection Controls",
            "The literature matrix aligns studies against researcher-defined dimensions. Cells can "
            "be edited and inspected so that extracted values remain connected to evidence and to "
            "the researcher’s corrections.",
            "image6.png",
            "Figure A3. Literature Review Matrix with editable structured evidence fields.",
        ),
        (
            "A.2.4 Final Research Outcomes",
            "The final stage consolidates the selected research topic, research question, study "
            "design, and proposed paper structure. These outputs remain downstream of the recorded "
            "search, screening, matrix, synthesis, and gap-analysis artifacts.",
            "image7.png",
            "Figure A4. Final Question workspace linking the selected outcome to study design and paper structure.",
        ),
        (
            "A.2.5 Peer-Review Access",
            "The peer-review interface exposes the project for independent inspection without "
            "transferring ownership of the researcher’s decisions. It supports review of the "
            "recorded process and artifacts rather than treating the final narrative as the only output.",
            "image8.png",
            "Figure A5. Peer-review access for inspecting workflow records and research artifacts.",
        ),
    ]
    for index, spec in enumerate(figure_specs):
        subsection, intro, filename, caption = spec
        add_figure(
            doc,
            subsection=subsection,
            intro=intro,
            image_path=ASSET_DIR / filename,
            caption=caption,
            final=index == len(figure_specs) - 1,
        )

    add_highlighted_paragraph(
        doc,
        "Source note: Screenshots are from the Matrica project interface supplied with the "
        "project materials and show illustrative interface states available at the time captured.",
        italic=True,
        size=8.5,
    )


def add_appendix_b(doc: Document) -> int:
    workbook = load_workbook(SOURCE_XLSX, data_only=True)
    sheet = workbook["Competitor Matrix"]
    header_map = {
        sheet.cell(2, column).value: column
        for column in range(1, sheet.max_column + 1)
        if sheet.cell(2, column).value
    }
    selected_rows = [
        row
        for row in range(3, sheet.max_row + 1)
        if sheet.cell(row, header_map["Similarity to Matrica"]).value == 5
        and get_value(sheet, row, header_map, "Product Name") != "Parsifal / ParsifalPro"
    ]
    if len(selected_rows) != 17:
        raise RuntimeError(f"Expected 17 selected products; found {len(selected_rows)}")

    names = [get_value(sheet, row, header_map, "Product Name") for row in selected_rows]
    urls = [get_value(sheet, row, header_map, "Official URL") for row in selected_rows]

    doc.add_page_break()
    add_heading(doc, "Appendix B. Research AI Systems with Similarity Score 5", 1)
    add_highlighted_paragraph(
        doc,
        "This appendix is derived directly from the ‘Competitor Matrix’ worksheet in "
        "Matrica_Competitor_Matrix_All_2026-08-26.xlsx. It includes every row whose "
        "‘Similarity to Matrica’ value is 5, except Parsifal / ParsifalPro, which was "
        "excluded because its primary interface language is not English. The resulting "
        "set contains 17 products: Matrica and 16 comparison systems. Entries describe "
        "the workbook snapshot and may change as products evolve.",
    )

    add_heading(doc, "B.1 Product Overview", 2)
    overview_rows = [
        [
            get_value(sheet, row, header_map, "Product Name"),
            get_value(sheet, row, header_map, "Product Type"),
            get_value(sheet, row, header_map, "Subject Focus"),
            get_value(sheet, row, header_map, "Closest Matrica Module"),
            get_value(sheet, row, header_map, "Short Description"),
        ]
        for row in selected_rows
    ]
    make_table(
        doc,
        ["Product", "Type", "Focus", "Closest Matrica module", "Workbook description"],
        overview_rows,
        [1.05, 0.65, 0.85, 1.35, 2.6],
        font_size=7.2,
        hyperlink_column=0,
        hyperlinks=urls,
    )
    add_highlighted_paragraph(
        doc,
        "Table B1. Product overview for the similarity-5 subset. Product names link to the "
        "official URLs recorded in the workbook.",
        style="Caption",
        italic=True,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        size=8.5,
    )

    landscape = doc.add_section(WD_SECTION.NEW_PAGE)
    landscape.orientation = WD_ORIENT.LANDSCAPE
    landscape.page_width = Inches(11)
    landscape.page_height = Inches(8.5)
    landscape.top_margin = Inches(0.55)
    landscape.bottom_margin = Inches(0.55)
    landscape.left_margin = Inches(0.55)
    landscape.right_margin = Inches(0.55)

    add_heading(doc, "B.2 Workflow Capability Comparison", 2)
    capability_fields = [
        "Direct Database Search",
        "AI Screening",
        "Data Extraction",
        "Literature Matrix",
        "Qualitative Synthesis",
        "Research Gap Identification",
        "Research Question Generation",
    ]
    capability_rows = [
        [get_value(sheet, row, header_map, "Product Name")]
        + [get_value(sheet, row, header_map, field) for field in capability_fields]
        for row in selected_rows
    ]
    make_table(
        doc,
        [
            "Product",
            "Direct search",
            "AI screening",
            "Data extraction",
            "Literature matrix",
            "Qualitative synthesis",
            "Gap identification",
            "RQ generation",
        ],
        capability_rows,
        [1.55, 1.18, 1.18, 1.18, 1.18, 1.28, 1.18, 1.22],
        font_size=6.7,
        hyperlink_column=0,
        hyperlinks=urls,
    )
    add_highlighted_paragraph(
        doc,
        "Table B2. Selected workflow capabilities exactly as coded in the workbook.",
        style="Caption",
        italic=True,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        size=8.2,
    )

    doc.add_page_break()
    add_heading(doc, "B.3 Provenance, Governance, and Reproducibility", 2)
    governance_fields = [
        "Evidence Provenance",
        "Human-in-the-loop",
        "Decision Journal",
        "Version History",
        "Audit Trail",
        "Reproducibility Support",
        "Open Source",
    ]
    governance_rows = [
        [get_value(sheet, row, header_map, "Product Name")]
        + [get_value(sheet, row, header_map, field) for field in governance_fields]
        for row in selected_rows
    ]
    make_table(
        doc,
        [
            "Product",
            "Evidence provenance",
            "Human in loop",
            "Decision journal",
            "Version history",
            "Audit trail",
            "Reproducibility",
            "Open source",
        ],
        governance_rows,
        [1.55, 1.25, 1.18, 1.2, 1.18, 1.08, 1.28, 1.18],
        font_size=6.7,
        hyperlink_column=0,
        hyperlinks=urls,
    )
    add_highlighted_paragraph(
        doc,
        "Table B3. Provenance, governance, and reproducibility fields exactly as coded in the workbook.",
        style="Caption",
        italic=True,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        size=8.2,
    )
    add_highlighted_paragraph(
        doc,
        "Source: Matrica_Competitor_Matrix_All_2026-08-26.xlsx, ‘Competitor Matrix’ sheet, "
        "filtered to Similarity to Matrica = 5 and excluding Parsifal / ParsifalPro.",
        italic=True,
        size=8.2,
    )
    return len(selected_rows)


def validate_original_content(source: Document, output: Document) -> None:
    if len(output.paragraphs) <= len(source.paragraphs):
        raise RuntimeError("Appendix paragraphs were not added")
    for index, source_paragraph in enumerate(source.paragraphs):
        output_paragraph = output.paragraphs[index]
        if source_paragraph.text != output_paragraph.text:
            raise RuntimeError(f"Original paragraph text changed at index {index}")
        if source_paragraph.style.style_id != output_paragraph.style.style_id:
            raise RuntimeError(f"Original paragraph style changed at index {index}")
        if source_paragraph._p.xml != output_paragraph._p.xml:
            raise RuntimeError(f"Original paragraph XML changed at index {index}")

    if len(output.tables) != len(source.tables) + 4:
        raise RuntimeError(
            f"Expected {len(source.tables) + 4} tables; found {len(output.tables)}"
        )
    for table_index, source_table in enumerate(source.tables):
        output_table = output.tables[table_index]
        if source_table._tbl.xml != output_table._tbl.xml:
            raise RuntimeError(f"Original table changed at index {table_index}")


def count_highlighted_runs(doc: Document, starting_paragraph: int) -> int:
    count = 0
    for paragraph in doc.paragraphs[starting_paragraph:]:
        count += sum(run.font.highlight_color == YELLOW for run in paragraph.runs)
    for table in doc.tables[3:]:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    count += sum(run.font.highlight_color == YELLOW for run in paragraph.runs)
                    for hyperlink in paragraph._p.findall(qn("w:hyperlink")):
                        count += len(hyperlink.findall(".//" + qn("w:highlight")))
    return count


def main() -> None:
    for required in [SOURCE_DOCX, SOURCE_XLSX, *(ASSET_DIR / f"image{i}.png" for i in range(4, 9))]:
        if not required.exists():
            raise FileNotFoundError(required)

    source = Document(SOURCE_DOCX)
    source_paragraph_count = len(source.paragraphs)
    doc = Document(SOURCE_DOCX)

    add_appendix_a(doc)
    selected_count = add_appendix_b(doc)

    OUTPUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_DOCX)

    output = Document(OUTPUT_DOCX)
    validate_original_content(source, output)
    highlighted_runs = count_highlighted_runs(output, source_paragraph_count)
    if highlighted_runs < 250:
        raise RuntimeError(f"Unexpectedly low highlighted-run count: {highlighted_runs}")

    print(f"output={OUTPUT_DOCX}")
    print(f"original_paragraphs_preserved={source_paragraph_count}")
    print(f"original_tables_preserved={len(source.tables)}")
    print(f"selected_products={selected_count}")
    print(f"appended_tables={len(output.tables) - len(source.tables)}")
    print(f"appended_images={len(output.inline_shapes)}")
    print(f"highlighted_runs={highlighted_runs}")


if __name__ == "__main__":
    main()
