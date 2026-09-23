from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Pt


SOURCE = Path("/Volumes/T9/Research/Research-AI/paper/Draft V3.docx")
OUTPUT = Path("/Volumes/T9/TPP/topic-generator/output/Draft V3_revised_highlighted.docx")


def find_paragraph(doc, exact_text):
    matches = [p for p in doc.paragraphs if p.text == exact_text]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph for {exact_text!r}, found {len(matches)}")
    return matches[0]


def base_rpr(paragraph):
    for run in paragraph.runs:
        if run._r.rPr is not None:
            return deepcopy(run._r.rPr)
    return None


def clear_runs(paragraph):
    for child in list(paragraph._p):
        if child.tag.endswith("}pPr"):
            continue
        paragraph._p.remove(child)


def add_run(paragraph, text, highlighted=False, rpr=None):
    run = paragraph.add_run(text)
    if rpr is not None:
        existing = run._r.rPr
        if existing is not None:
            run._r.remove(existing)
        run._r.insert(0, deepcopy(rpr))
    if highlighted:
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return run


def set_segments(paragraph, segments):
    rpr = base_rpr(paragraph)
    clear_runs(paragraph)
    for text, highlighted in segments:
        add_run(paragraph, text, highlighted, rpr)


def append_highlighted(paragraph, text):
    rpr = base_rpr(paragraph)
    add_run(paragraph, text, True, rpr)


def insert_heading_after(doc, paragraph, heading_text, model_heading):
    new_p = doc.add_paragraph()
    if new_p._p.pPr is not None:
        new_p._p.remove(new_p._p.pPr)
    if model_heading._p.pPr is not None:
        new_p._p.insert(0, deepcopy(model_heading._p.pPr))
    add_run(new_p, heading_text, True, base_rpr(model_heading))
    paragraph._p.addnext(new_p._p)
    return new_p


def main():
    doc = Document(SOURCE)

    p = find_paragraph(
        doc,
        "Research AI systems differ in the type of research activity they primarily support and in the form in which they organise their outputs. Some systems are centred on source-grounded document analysis, some on answering research questions through web or academic search, and others on structured evidence-review tasks. To clarify Matrica’s position, Table 1 compares four representative systems with Matrica in terms of their primary emphasis and the main structure through which research work is organised. The comparison is conceptual rather than a benchmark of product quality or accuracy.",
    )
    append_highlighted(
        p,
        " Appendix B complements this conceptual comparison with a dated, workbook-coded view of the 16 comparison systems assigned the highest similarity score in the accompanying competitor matrix.",
    )

    p = find_paragraph(
        doc,
        "The comparison positions Matrica as workflow infrastructure rather than as another search engine, chatbot, or answer-generation system. Its principal unit is the evolving literature-review project rather than an individual conversation, answer, document collection, or evidence table. Matrica does not claim that its individual functions are unique. Its contribution lies in connecting these functions within one structured process and retaining selected records that allow the development of the review and the researcher’s participation to be examined.",
    )
    append_highlighted(
        p,
        " The Appendix B classifications provide descriptive landscape context only; they reproduce the workbook snapshot and are not independently validated measures of product quality or comparative performance.",
    )

    p = find_paragraph(
        doc,
        "To examine this contribution, this paper distinguishes three related forms of provenance. Evidence provenance identifies the source material supporting a value or claim; process provenance identifies the activities through which that value or claim was produced; and contribution provenance identifies how the researcher evaluated, changed, or used the available evidence and AI assistance. This distinction builds on general provenance models that relate entities, activities, and responsible agents [41], while contribution provenance is introduced here as a Matrica-specific extension for examining human participation in AI assisted research. Where feasible, these records should be captured during the research process because later descriptions may not accurately reconstruct the evidence, system output, decision, and justification available at the time.These principles provide the conceptual basis for Matrica’s human-control and review mechanisms. Matrica records selected consequential decisions and justifications alongside AI-assisted research artifacts so that an independent reviewer can examine how the researcher participated in producing the review. The reviewer evaluates the recorded reasoning and interventions without replacing the researcher’s decisions. The specific mechanisms through which Matrica implements this approach are described in Section 5.",
    )
    set_segments(
        p,
        [
            (
                "To examine this contribution, this paper distinguishes three related forms of provenance. Evidence provenance identifies the source material supporting a value or claim; process provenance identifies the activities through which that value or claim was produced; and contribution provenance identifies how the researcher evaluated, changed, or used the available evidence and AI assistance. This distinction builds on general provenance models that relate entities, activities, and responsible agents [41], while contribution provenance is introduced here as a Matrica-specific extension for examining human participation in AI assisted research. Where feasible, these records should be captured during the research process because later descriptions may not accurately reconstruct the evidence, system output, decision, and justification available at the time. ",
                False,
            ),
            (
                "These principles provide the conceptual basis for Matrica’s human-control and review mechanisms.",
                True,
            ),
            (
                " Matrica records selected consequential decisions and justifications alongside AI-assisted research artifacts so that an independent reviewer can examine how the researcher participated in producing the review. The reviewer evaluates the recorded reasoning and interventions without replacing the researcher’s decisions. The specific mechanisms through which Matrica implements this approach are described in Section 5.",
                False,
            ),
            (
                " Appendix A provides corresponding interface views for the principal controls discussed in Sections 5 and 6.",
                True,
            ),
        ],
    )

    p = find_paragraph(
        doc,
        "This gap motivates Matrica’s positioning as a provenance-aware, matrix-centred, and human-controlled workflow. The present paper describes the framework and implements artifact, future empirical work should evaluate output reliability, error detection, appropriate reliance, recovery from change, and the visibility of substantive human judgment.",
    )
    set_segments(
        p,
        [
            (
                "This gap motivates Matrica’s positioning as a provenance-aware, matrix-centred, and human-controlled workflow. ",
                False,
            ),
            (
                "The present paper describes the framework and its implemented artifact; future empirical work should evaluate output reliability, error detection, appropriate reliance, recovery from change, and the visibility of substantive human judgment.",
                True,
            ),
        ],
    )

    p = find_paragraph(
        doc,
        "> AI can also assist with evidence organisation, matrix completion, and preliminary synthesis [25, 27, 32]. These activities require comparison across studies and interpretation of incomplete or conflicting findings, and their reliability is less easily represented by a single classification score. The system may generate candidate themes, gaps, or research questions, but the plausibility of such outputs does not establish that they follow from the underlying evidence. Downstream proposals should therefore remain connected to the intermediate evidence and decisions from which they were produced and should be treated as researcher-reviewed outputs rather than established findings.",
    )
    set_segments(
        p,
        [
            (
                "AI can also assist with evidence organisation, matrix completion, and preliminary synthesis [25, 27, 32].",
                True,
            ),
            (
                " These activities require comparison across studies and interpretation of incomplete or conflicting findings, and their reliability is less easily represented by a single classification score. The system may generate candidate themes, gaps, or research questions, but the plausibility of such outputs does not establish that they follow from the underlying evidence. Downstream proposals should therefore remain connected to the intermediate evidence and decisions from which they were produced and should be treated as researcher-reviewed outputs rather than established findings.",
                False,
            ),
        ],
    )

    p = find_paragraph(
        doc,
        "> Within the DSRM process, this study covers problem identification, the definition of solution objectives, and the design and development of an executable Matrica prototype. The study does not include a documented application of Matrica to a completed literature review, an adjudicated benchmark, or an observed user study. Accordingly, this paper reports the implemented artifact and defines the boundary of the claims that can currently be made; demonstration in use and empirical evaluation remain future work.",
    )
    set_segments(
        p,
        [
            (
                "Within the DSRM process, this study covers problem identification, the definition of solution objectives, and the design and development of an executable Matrica prototype.",
                True,
            ),
            (
                " The study does not include a documented application of Matrica to a completed literature review, an adjudicated benchmark, or an observed user study. Accordingly, this paper reports the implemented artifact and defines the boundary of the claims that can currently be made; demonstration in use and empirical evaluation remain future work.",
                False,
            ),
        ],
    )

    p = find_paragraph(
        doc,
        "The sequence converts a broad interest into structured intent, an analytic schema, a retrieval and screening record, comparable evidence, and a defensible question. Figure 1 shows the implemented path.",
    )
    append_highlighted(
        p,
        " Appendix A maps these stages to their principal interface controls and provides selected workflow views.",
    )

    figure_1_caption = find_paragraph(
        doc,
        "Figure 1. Implemented nine-stage Matrica workflow. AI proposes or classifies within stages; the researcher defines, selects, inspects, edits, justifies, and confirms the research state.",
    )
    figure_1_caption.paragraph_format.space_before = Pt(24)

    p = find_paragraph(
        doc,
        "The common interaction pattern is proposed inspect decide record. AI services can generate topics, queries, classifications, matrix content, synthesis blocks, gaps, and final artifacts. The user may select, edit, skip, replace, or request regeneration. The application then persists the resulting project state and, for selected consequential actions, adds a journal entry. This pattern makes a distinction between a model suggestion and the current user-controlled state.",
    )
    set_segments(
        p,
        [
            ("The common interaction pattern is propose–inspect–decide–record.", True),
            (
                " AI services can generate topics, queries, classifications, matrix content, synthesis blocks, gaps, and final artifacts. The user may select, edit, skip, replace, or request regeneration. The application then persists the resulting project state and, for selected consequential actions, adds a journal entry. This pattern distinguishes a model suggestion from the current user-controlled state.",
                False,
            ),
        ],
    )

    p = find_paragraph(
        doc,
        "Stage 1 uses a conversational assistant to elicit five canonical fields including, domain, population, focus, method, and review type. The fields are stored separately from the chat transcript and passed forward through a serialized research context. The user can edit any field directly. Profile edits create audit entries containing the previous and new values, allowing later review of how the project's framing changed.",
    )
    set_segments(
        p,
        [
            (
                "Stage 1 uses a conversational assistant to elicit five canonical fields: domain, population, focus, method, and review type.",
                True,
            ),
            (
                " The fields are stored separately from the chat transcript and passed forward through a serialized research context. The user can edit any field directly. Profile edits create audit entries containing the previous and new values, allowing later review of how the project's framing changed.",
                False,
            ),
            (" Figure A1 shows the Research Profile within the persistent workflow shell.", True),
        ],
    )

    p = find_paragraph(
        doc,
        "The AI classifier is instructed to use the supplied criteria and can fetch a real abstract when the local record is insufficient. An uncertain AI result is routed into the review list as pending rather than silently excluded. Users can override AI decisions, clear AI made screening decisions while leaving human decisions intact, add notes, and move papers among states. This actor distinction is one of the clearest operationalizations of human contribution. The final inclusion set can be inspected together with who made each current decision and why.",
    )
    set_segments(
        p,
        [
            (
                "The AI classifier is instructed to use the supplied criteria and can fetch a real abstract when the local record is insufficient. An uncertain AI result is routed into the review list as pending rather than silently excluded. Users can override AI decisions, clear ",
                False,
            ),
            ("AI-made", True),
            (
                " screening decisions while leaving human decisions intact, add notes, and move papers among states. This actor distinction is one of the clearest operationalizations of human contribution. The final inclusion set can be inspected together with who made each current decision and why.",
                False,
            ),
            (" Figure A2 illustrates these protocol controls and decision states.", True),
        ],
    )

    p = find_paragraph(
        doc,
        "The current interface displays cell content and provides an on-demand evidence and reasoning popover. Opening that view is journaled as an AI explanation and evidence interaction. However, the present autofill path initializes the source reference field as empty, and the popover's explanation and sources are generated on demand rather than stored as immutable source spans. The implementation therefore supports inspection and records that inspection occurred, but it should not yet be described as complete passage level provenance or as a fully implemented verification checkbox workflow. Section 10 treats this as a priority limitation.",
    )
    set_segments(
        p,
        [
            (
                "The current interface displays cell content and provides an on-demand evidence and reasoning popover. Opening that view is journaled as an AI explanation and evidence interaction. However, the present autofill path initializes the source reference field as empty, and the popover's explanation and sources are generated on demand rather than stored as immutable source spans. The implementation therefore supports inspection and records that inspection occurred, but it should not yet be described as complete passage level provenance or as a fully implemented verification checkbox workflow. ",
                False,
            ),
            ("Section 8 treats this as a priority limitation.", True),
            (" Figure A3 illustrates the matrix and its inspection surface without extending that provenance claim.", True),
        ],
    )

    p = find_paragraph(
        doc,
        "This stage demonstrates why upstream provenance matters. A polished question is not independently valid; it inherits the scope, search, selection, and extraction decisions made earlier. Matrica retains those upstream objects in the same project snapshot so that a reviewer can inspect the route to the final choice.",
    )
    append_highlighted(p, " Figure A4 shows the corresponding final-question workspace.")

    p = find_paragraph(
        doc,
        "Taken together, these mechanisms make human contribution more visible than a final report alone. They still require interpretation. A long journal can document activity without demonstrating good reasoning, while a concise intervention may be decisive. Future evaluation should therefore assess the appropriateness and consequences of human actions rather than merely their frequency.6. System Architecture and Implementation",
    )
    set_segments(
        p,
        [
            (
                "Taken together, these mechanisms make human contribution more visible than a final report alone. They still require interpretation. A long journal can document activity without demonstrating good reasoning, while a concise intervention may be decisive. Future evaluation should therefore assess the appropriateness and consequences of human actions rather than merely their frequency.",
                False,
            )
        ],
    )
    model_heading = find_paragraph(doc, "7. Discussion")
    insert_heading_after(doc, p, "6. System Architecture and Implementation", model_heading)

    p = find_paragraph(
        doc,
        "The peer review subsystem creates revocable bearer token links. The reviewer receives a redacted project snapshot organized by all nine stages and can comment on each stage's result and on recorded user or AI evidence operations. Reviewers cannot include, exclude, approve, or replace the user's decisions. Comments autosave, can be submitted, and are displayed to the owner with invitation status.",
    )
    append_highlighted(p, " Figure A5 shows the corresponding peer-review access interface.")

    p = find_paragraph(
        doc,
        "Table 3 summarizes the inspected implementation and verification boundary. All automated checks were executed against the repository revision identified in Section 4.3.",
    )
    set_segments(
        p,
        [
            ("Table 3 summarizes the inspected implementation and verification boundary. All automated checks were executed against the repository revision identified ", False),
            ("under Software and Data Availability", True),
            (".", False),
        ],
    )

    p = find_paragraph(
        doc,
        "The design process produced seven requirements and an implemented nine-stage artifact. Matrica combines shared structured state, stage-specific services, actor-marked screening, evidence matrices, assistant-action confirmation, selected decision records, snapshots, peer-review access, analytics, and export. These outcomes demonstrate the implementation of the proposed workflow mechanisms; they do not establish improvements in review accuracy, usability, appropriate reliance, or scholarly contribution, which require future application and empirical evaluation.",
    )
    append_highlighted(
        p,
        " Appendix A illustrates where selected mechanisms appear in the interface, while Appendix B provides the dated system-landscape context for Matrica's positioning.",
    )

    p = find_paragraph(
        doc,
        "Four implications follow. First, Research AI should represent the researcher's epistemic work, not only model output. Second, provenance should be layered: source, transformation, and contribution provenance answer different questions. Third, evaluation should test the workflow in which outputs are used, because a model's task accuracy does not reveal whether users notice its failures. Fourth, implementation papers should distinguish an intended mechanism from a verified feature and a verified feature from an empirically beneficial outcome.",
    )
    append_highlighted(
        p,
        " The workbook-derived classifications in Appendix B further illustrate the need to date comparative system maps and state their coding and validation boundaries.",
    )

    p = find_paragraph(
        doc,
        "The first limitation is the absence of summative empirical evaluation. The repository and its automated tests were inspected, but no adjudicated benchmark or participant study was available. Passing tests demonstrate implementation integrity within their coverage; they do not demonstrate review validity, usability, or improved human reasoning.",
    )
    append_highlighted(
        p,
        " Accordingly, the screenshots in Appendix A document interface states rather than observed use, and the comparisons in Appendix B reproduce workbook coding rather than benchmark results.",
    )

    p = find_paragraph(
        doc,
        "The next step is to apply Matrica in documented literature-review projects and evaluate how the workflow performs in practice. Future work should examine whether it supports reliable review outcomes, clearer researcher accountability, effective recovery from changes, and meaningful independent review. It should also strengthen source-level evidence provenance and the tracking of dependencies among research artifacts. These findings will guide the continued development of Matrica as a framework for accountable human–AI collaboration in literature reviews.",
    )
    set_segments(
        p,
        [
            (
                "The next step is to apply Matrica in documented literature-review projects and evaluate how the workflow performs in practice. Future work should examine whether it supports reliable review outcomes, clearer researcher accountability, effective recovery from changes, and meaningful independent review. It should also strengthen source-level evidence provenance and the tracking of dependencies among research artifacts. ",
                False,
            ),
            (
                "Appendix A makes selected artifact mechanisms visible, while Appendix B supplies descriptive system context; neither substitutes for empirical evaluation.",
                True,
            ),
            (
                " These findings will guide the continued development of Matrica as a framework for accountable human–AI collaboration in literature reviews.",
                False,
            ),
        ],
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
