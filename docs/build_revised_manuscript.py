from __future__ import annotations

import re
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(
    "/Users/jingxiansu/.codex/attachments/"
    "b923a82b-283f-49fb-930e-c8f814656338/pasted-text.txt"
)
OUTPUT = ROOT / "docs" / "Matrica_Revised_Manuscript.docx"
ASSET_DIR = ROOT / "docs" / "manuscript_assets"

BLUE = "#2E5D7B"
DARK_BLUE = "#173B52"
LIGHT_BLUE = "#EAF2F7"
LIGHT_GRAY = "#F4F6F9"
MID_GRAY = "#66737C"
WHITE = "#FFFFFF"
BLACK = "#1F2529"
GOLD = "#B6852D"

FONT_REGULAR = "/System/Library/Fonts/Supplemental/Arial.ttf"
FONT_BOLD = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"


def extract(text: str, start: str, end: str) -> str:
    i = text.index(start)
    j = text.index(end, i)
    return text[i:j].strip()


source_text = SOURCE.read_text(encoding="utf-8").replace("\r\n", "\n")

related_work_31_33 = extract(
    source_text,
    "3.1 Structured Literature Review and Evidence Synthesis",
    "3.4 Human–AI Interaction, Trust, and Provenance",
)
related_work_31_33 = related_work_31_33.replace(
    "stale-state propagation", "step-level stale-state signalling"
).replace(
    "dependency-aware change handling", "step-level stale-state handling"
)

evaluation_72_76 = extract(
    source_text,
    "7.2 Baselines",
    "8. Design Results and Empirical Evidence Status",
)

references = extract(source_text, "References", "[44]")
ref44 = source_text[source_text.rindex("[44]"):].strip()
references = references + "\n" + ref44
references += (
    "\n[45] TPP Projects, \"Matrica,\" GitHub repository, 2026. "
    "Available: https://github.com/TPP-PROJECTS/Matrica. "
    "Implementation revision inspected: 9e30bc4, 26 July 2026."
)


front_matter = """1. Abstract
Artificial-intelligence assistance is becoming a routine part of literature-review work, particularly in searching, title-and-abstract screening, data extraction, and synthesis. For many review teams, the practical question is therefore no longer simply whether AI should be used, but how its use can remain inspectable, contestable, and accountable. Evaluating only the final answer conceals the researcher's intellectual work: framing the problem, defining comparison dimensions, choosing sources and queries, setting eligibility criteria, resolving uncertainty, correcting generated content, and deciding which evidence may support a claim. This paper presents Matrica, an implemented open-source, provenance-aware workflow designed to make both AI participation and human contribution visible across a nine-stage path from research intent to a final research question and study design. Following Design Science Research Methodology, the system operationalizes seven requirements: structured intent, explicit process state, transparent selection, comparable evidence, observable human agency, recoverable change, and evidence-aligned outcomes. The implementation combines a typed React/TypeScript client, modular FastAPI services, OpenAlex and PubMed retrieval, protocol-based screening, literature and system matrices, a decision journal, human-confirmed assistant actions, local and server-side snapshots, peer process review, workflow analytics, and multi-format export. Verification against repository revision 9e30bc4 produced 77 passing backend tests, 42 passing frontend tests, a clean TypeScript check, and a successful production build. These results establish implementation and technical conformance, not improved review accuracy, efficiency, usability, or trust. A preregistration-ready protocol is therefore retained to evaluate retrieval, screening, extraction, workload, error detection, appropriate reliance, and the substantive character of human contribution. Matrica's principal contribution is not another automatic answer generator; it is an operational model for examining how people use AI, how they judge its outputs, and how their decisions shape a research product.
Keywords: Research AI; systematic literature review; literature review; human–AI collaboration; human contribution; provenance; evidence synthesis; appropriate reliance; Design Science Research

2. Introduction
2.1 Context and Motivation
Generative AI has moved rapidly from an experimental add-on to a recurring component of academic search and evidence-synthesis practice. A recent rapid review identified 115 empirical studies of generative AI across systematic-review tasks, with the strongest evidence concentrated in title-and-abstract screening and data extraction [1]. A separate LLM-assisted systematic review mapped 172 eligible publications and documented the use of language models not only as objects of evaluation but also within review production itself [2, 3]. This diffusion does not mean that every researcher must use AI or that conventional review methods are obsolete. It does mean that a blanket assumption of an AI-free workflow is increasingly unrealistic in settings where AI functions are embedded in research assistants, discovery systems, writing environments, and specialist review tools.

Systematic and structured literature reviews are especially important test cases. They require large numbers of linked decisions: defining scope, translating concepts into queries, searching multiple sources, screening records, extracting comparable evidence, interpreting contradictions, and converting a synthesis into a defensible gap or question [5, 8, 9]. Screening alone can dominate reviewer effort, while low inclusion prevalence and continually expanding literatures create strong incentives to use automated triage and extraction [6, 8, 10]. The attraction of AI is therefore not reducible to novelty or convenience. It responds to a real mismatch between the scale of available literature and the time available for careful human review.

Once AI assistance enters this chain, however, the central methodological problem changes. The relevant question is not merely whether a model can generate a fluent summary or reproduce a benchmark label. Researchers must know what evidence the system received, what it proposed, where uncertainty or failure occurred, whether a person inspected the basis of the proposal, what the person changed, and how that intervention altered downstream work. A final research question may contain no visible trace of the human decisions that made it possible, even though those decisions constitute much of the project's epistemic contribution.

Matrica was developed in response to this governance and observability problem. Tool fragmentation remains relevant because moving data among chat, search, spreadsheet, and writing applications often strips away context. It is not, however, the sole or even primary motivation. The deeper problem is that AI use can become methodologically invisible: model assistance is hard to reconstruct, human judgment is compressed into a final accept/reject state, and the reliability of the final result is inferred from presentation quality rather than from the quality of the process that produced it. Matrica therefore treats the research workflow itself as an object that must be represented, inspected, evaluated, and, where appropriate, reviewed by another person.

2.2 Problem Statement
Current Research AI presents at least six connected risks. First, process opacity makes it difficult to recover why a query was chosen, why a paper was excluded, or which evidence informed a matrix value [5, 11, 19]. Second, model performance is conditional on prompt wording, criteria, evidence access, model version, and task structure; results that are reliable for explicit study settings may be weaker for nuanced outcomes or contextual judgments [20-24]. Third, errors propagate: incomplete retrieval constrains screening, false exclusions constrain extraction, and weak extraction constrains synthesis and gap claims [16, 23, 24].

Fourth, fluent and authoritative language can encourage automation bias. An explanation or confidence score may look like evidence even when it is uncalibrated or generated after the underlying value [9, 15, 25]. Fifth, missingness is heterogeneous. A blank may mean that a source did not report an item, that full text was unavailable, that a request failed, or that a model refused to respond; these states should not be collapsed into one absence [19, 26]. Sixth, the researcher's work is underrepresented. A human-in-the-loop label often reveals only that a person clicked approve, not whether the person defined the criterion, inspected the source, corrected the value, articulated a reason, or reconsidered an upstream assumption.

Fragmentation amplifies all six risks because transfers among tools detach outputs from the context in which they were produced [4, 17, 18]. Yet integration by itself is insufficient. A single interface can still hide model uncertainty, overwrite earlier states, or reduce human participation to ceremonial confirmation. The design problem is therefore to integrate the workflow while preserving meaningful distinctions among source evidence, AI suggestion, human decision, and downstream consequence.

This framing also changes how reliability should be evaluated. Output accuracy remains essential, especially sensitivity in screening and source-supported accuracy in extraction. It is nevertheless only one layer. A defensible evaluation must also ask whether users detect planted errors, challenge incorrect suggestions, preserve relevant uncertainty, document their reasons, recover after upstream changes, and avoid accepting unsupported synthesis because it is well written. Human contribution is not a substitute for technical accuracy, but neither is it captured by technical accuracy alone.

2.3 Research Gap
Most empirical studies of AI-assisted reviews evaluate bounded tasks such as title-and-abstract screening, full-text screening, data extraction, or risk-of-bias assessment [1, 2]. Increasingly broad systems such as AiReview, otto-SR, and AutoSynthesis demonstrate multi-stage or end-to-end capability, but their principal evaluations emphasize task performance, reproduction of completed reviews, workload reduction, or agreement with expert outputs [6, 7, 27]. These are valuable outcomes, but they provide less detail about how a researcher's interventions are represented across stages.

Human-in-the-loop proposals commonly recommend verification, conservative thresholds, structured prompts, and expert adjudication [11, 28, 29]. The loop is often modeled as a safeguard around model output rather than as a substantive part of knowledge production. As a result, the researcher's contribution may be visible only at points of disagreement, even though people also contribute by defining the analytic schema, selecting which sources and systems are relevant, interpreting ambiguous evidence, deciding when missingness reflects a search deficiency, and determining whether a proposed gap is consequential.

The resulting research gap is at the workflow and contribution levels. Prior work offers growing evidence that individual tasks can be assisted, but less evidence on systems that (1) carry structured research state across the full review, (2) distinguish AI-originated and human-originated changes, (3) make inspection and correction observable, (4) preserve reasons and reviewer commentary, and (5) support evaluation of whether human involvement was substantive rather than ritualized. Matrica addresses this gap by treating intermediate artifacts and selected interactions as inspectable research records rather than disposable steps toward a final report.

2.4 Aim, Research Questions, and Contributions
This study aims to design, implement, and specify the evaluation of Matrica as a human-controlled Research AI workflow that makes the production of a literature review traceable at both the evidence and contribution levels. The inquiry shifts attention from answer generation to the allocation of work and judgment between a model, a researcher, and, where used, a peer reviewer.

RQ1. Which workflow and governance requirements follow when AI assistance is treated as a normal but fallible component of systematic and structured literature-review practice?

RQ2. How can an implemented, provenance-aware, matrix-centered system represent AI suggestions, human decisions, and their downstream relationships from research intent to a final research question?

RQ3. How should the system be evaluated so that result reliability, human effort, appropriate reliance, and the substantive contribution of the researcher are measured without equating fluent output or frequent clicking with methodological quality?

The paper makes four contributions. First, it reframes the motivation for Research AI from overcoming tool fragmentation to governing increasingly routine AI participation in scholarly work. Second, it reports an implemented open-source, nine-stage artifact rather than a reference architecture alone. Third, it operationalizes human contribution through actor-marked screening decisions, edits, selections, justifications, evidence-view events, confirmed or skipped assistant actions, snapshots, and peer comments. Fourth, it proposes a falsifiable evaluation framework that separates implementation verification from future claims about accuracy, efficiency, usability, appropriate reliance, and human intellectual engagement.
"""


human_ai_related = """3.4 Human–AI Interaction, Appropriate Reliance, and Contribution
The objective of Research AI should be appropriate reliance rather than maximal trust. Humans and models fail differently: models can process and triage large document sets with speed and consistency, while researchers contribute contextual judgment, methodological interpretation, sensitivity to domain consequences, and accountability for the final review [10, 25, 31, 33]. Human participation is therefore not a residual error-correction layer. It is constitutive of the review's scope, coding frame, evidential thresholds, and interpretation.

Confidence labels and explanations can help users decide where to inspect, but they can also create false assurance when they are not calibrated against observed errors or linked to recoverable evidence [9, 21, 29]. Human–AI interaction guidance accordingly emphasizes making system capabilities clear, supporting efficient correction, showing relevant context, and allowing users to control or dismiss system actions [42]. In a review workflow, these principles require more than an editable text box. The interface must preserve what the AI proposed, what the person accepted or changed, and which source or criterion was available at the time.

This paper uses human contribution to mean the observable intellectual and procedural interventions through which a person shapes the research product. Examples include defining the Research Profile, approving or changing matrix dimensions, choosing databases or comparison systems, editing queries and filters, establishing eligibility criteria, including or excluding a paper with a reason, inspecting AI-provided evidence, choosing the synthesis dimensions that matter, rejecting a weak gap claim, and selecting a research question or study design. These actions do not all have equal scholarly value, and their count is not an authorship score. Their value depends on whether they are methodologically appropriate and whether they improve or constrain the resulting evidence claim.

Provenance and contribution records therefore serve different but related purposes. Evidence provenance asks where a value or claim came from. Process provenance asks which transformation produced it. Contribution provenance asks what the person did with the available evidence or AI proposal. A trustworthy workflow should capture all three at creation time where feasible, because a post hoc narrative cannot reliably reconstruct the exact prompt, source context, decision, or correction that shaped an artifact [5, 19].

Peer review adds an evaluative layer. A second person should be able to inspect not only the final matrix or question but also the recorded decisions and justifications that led there. Matrica therefore treats peer comments as assessments of the user's critical-thinking process rather than as silent replacements of the user's decisions. This distinction preserves accountability while making the quality of human involvement open to scrutiny.

3.5 Summary of the Literature Gap
The literature supports targeted use of generative AI for selected review tasks, especially screening and structured extraction, but it does not support indiscriminate end-to-end autonomy [1, 16, 23, 34]. Performance depends on evidence access, criteria, prompts, models, thresholds, and adjudication. At the same time, AI support is sufficiently widespread that responsible design cannot rely on prohibition or nondisclosure as its main safeguard.

The clearer opportunity is to build systems that expose how evidence and decisions move across stages and how human intervention changes the research state. Matrica is consequently positioned as a provenance-aware, matrix-centered, human-controlled workflow. Its contribution is evaluated through implementation coverage, traceability, recoverability, error detection, appropriate reliance, and the visibility of human judgment—not through fluency or the mere presence of a human approval gate.
"""


methodology = """4. Research Methodology
4.1 Design Science Research and Study Status
The study follows the six activities of Design Science Research Methodology (DSRM): problem identification and motivation; definition of solution objectives; design and development; demonstration; evaluation; and communication [40]. DSRM is appropriate because the research problem concerns both an observed weakness in current practice and the construction of a socio-technical artifact intended to address it. The artifact includes the workflow, state representation, interface rules, provenance records, reviewer pathway, data services, and evaluation boundary.

Problem identification is grounded in evidence that AI performance varies across retrieval, screening, extraction, and synthesis, while human oversight is necessary but often underspecified [1, 9, 16, 20-24]. Solution objectives translate those limitations into observable obligations. Variable retrieval motivates query and coverage records; false exclusion risk motivates uncertainty and reason-linked screening; uneven extraction motivates evidence metadata and inspection; automation bias motivates confirmation and correction; invisible human labor motivates a decision journal and peer process review; and upstream revision motivates recoverable state plus explicit staleness cues.

Unlike the earlier manuscript version, design and development are now reported as completed for an executable research prototype. Demonstration at this stage consists of end-to-end workflow availability, source-code inspection, automated tests, type checking, and production-build verification. Summative empirical evaluation has not been completed. The current DSRM status is therefore: problem identification, objectives, design, development, technical demonstration, and communication completed; comparative technical benchmarking and participant evaluation specified but pending.

4.2 Requirement Derivation
Seven requirements were derived by translating recurring methodological and interaction failures into system behavior. The requirements deliberately combine functional needs with governance needs because an integrated workflow can still be untrustworthy if it hides origin, correction, or uncertainty.

Table 2. Literature-derived design requirements
ID	Requirement	Operationalization
R1	Structured intent	Capture the domain, population, focus, method, review type, and user-defined boundaries in a reusable Research Profile.
R2	Explicit process state	Persist the outputs of nine workflow stages as addressable fields in a shared project snapshot.
R3	Transparent selection	Expose queries, filters, screening criteria, include/exclude/uncertain states, reasons, and decision actor where implemented.
R4	Comparable evidence	Organize papers and optional systems against a user-controlled matrix schema with cell-level evidence metadata.
R5	Observable human agency	Record selected user edits, choices, justifications, evidence-view events, assistant-action outcomes, and peer comments.
R6	Recoverable change	Provide undo/redo, autosave, restorable project snapshots, and visible step-level staleness after confirmed back-navigation.
R7	Evidence-aligned outcomes	Carry structured matrix and synthesis state into gap analysis, final topics, research questions, study design, and export.

R5 is central to the revised positioning. Human control is not established merely because editing is technically possible. The system must make at least some consequential interventions observable and available for review. At the same time, the requirement does not imply that every mouse movement should be logged. Excessive instrumentation can create surveillance, privacy, and interpretive problems. Matrica therefore prioritizes selected decisions with methodological meaning.

4.3 Artifact Inspection and Evidence Boundary
The implementation description was verified against the open-source repository at revision 9e30bc4, dated 26 July 2026 [45]. Inspection covered the typed client state, reducer actions, nine screen routes, stage-specific API routers, retrieval and screening logic, matrix ingestion and extraction, persistence, analytics, exports, peer review, and automated tests. Verification was rerun on 3 August 2026. The backend test suite completed 77 tests successfully; the frontend completed 42 tests across five test files; TypeScript checking completed without error; and the Vite production build succeeded.

These observations support claims about what is implemented and whether the available automated checks pass. They do not establish retrieval recall, screening sensitivity, extraction accuracy, efficiency, usability, workload, trust calibration, or improved scholarly quality. No benchmark corpus, adjudicated output dataset, completed SUS or NASA-TLX instruments, or participant interview dataset was supplied for this paper. The empirical claims in Section 7 therefore remain hypotheses.

The artifact is described as PRISMA-informed rather than PRISMA-compliant. PRISMA 2020 is a reporting guideline, not a certification of review conduct or software quality [41]. Similarly, provenance-aware does not mean that every generated cell already contains immutable passage-level provenance. The current implementation stores structured cell metadata and offers an on-demand evidence-and-reasoning view, but the generated source explanation is not equivalent to a persistently stored source span. This limitation is made explicit rather than hidden behind the provenance label.
"""


framework = """5. The Implemented Matrica Workflow
5.1 Nine-Stage Workflow and Control Pattern
Matrica implements an ordered but revisable sequence of nine stages: (1) Chat and Define Topic, (2) Matrix Initial, (3) System Review, (4) Explore and Search, (5) Select Paper, (6) Matrix Table, (7) Synthesis, (8) Gap and Statement, and (9) Final Question. The sequence converts a broad interest into structured intent, an analytic schema, a retrieval and screening record, comparable evidence, and a defensible question. Figure 1 shows the implemented path.

[[FIGURE_1]]
Figure 1. Implemented nine-stage Matrica workflow. AI proposes or classifies within stages; the researcher defines, selects, inspects, edits, justifies, and confirms the research state.

The common interaction pattern is propose–inspect–decide–record. AI services can generate topics, queries, classifications, matrix content, synthesis blocks, gaps, and final artifacts. The user may select, edit, skip, replace, or request regeneration. The application then persists the resulting project state and, for selected consequential actions, adds a journal entry. This pattern makes a distinction between a model suggestion and the current user-controlled state.

5.2 Research Profile
Stage 1 uses a conversational assistant to elicit five canonical fields: domain, population, focus, method, and review type. The fields are stored separately from the chat transcript and passed forward through a serialized research context. The user can edit any field directly. Profile edits create audit entries containing the previous and new values, allowing later review of how the project's framing changed.

The limited field set is a pragmatic implementation choice rather than a universal ontology of research intent. It provides stable context for downstream prompts but does not replace a registered review protocol. More detailed eligibility boundaries, dates, languages, and search limits are represented in later stage-specific structures.

5.3 Matrix-First Schema Definition
Stage 2 asks the user to define the shape of the analysis before evidence extraction. The system recommends literature-matrix columns and, when enabled, system-comparison columns. The user can add, remove, and regenerate suggestions, provide optional reasons for schema changes, and keep the system matrix disabled when product or platform comparison is irrelevant.

The current implementation persists the active columns, recommendations, and column descriptions in project snapshots. It also records selected schema changes in the decision journal. It does not yet maintain a separate immutable schema-version identifier on each matrix cell. Accordingly, recoverability is provided through snapshots and undo/redo rather than a fully normalized version graph.

5.4 System Review
Stage 3 recommends databases, AI tools, or other systems relevant to the Research Profile. Each item contains a name, description, category, vendor, URL where available, AI-origin marker, and selection state. Users can select recommendations or add a custom system. These choices contribute to the optional system matrix and reveal which external systems the researcher considered relevant to the comparison.

The stage is analytically distinct from scholarly retrieval. A database, an AI assistant, and a review platform organize different research objects and should not be treated as interchangeable evidence sources. Recording the user's selection helps explain why a later comparison matrix contains a particular system set.

5.5 Topic, Query, and Search Construction
Stage 4 generates topic candidates grouped by population, design, risk, and review types. Cards retain a rationale, tags, confidence level, and whether the suggestion directly matched the profile or was inferred. Users select, edit, or add topics. Selected topics inform query generation, which records query text, target databases, quality score, priority, result count, and a stale flag.

Search filters include year range, open-access-only status, preprint exclusion, and language. User changes to filters and selected query edits can be journaled. The system therefore exposes important upstream choices that would otherwise disappear behind a search button. Current query history is represented through project snapshots and journal entries rather than a separate per-query version table.

5.6 Retrieval, Citation Chasing, and Deduplication
The backend searches OpenAlex and PubMed through source-specific adapters. PubMed abstracts are recovered through E-utilities, while OpenAlex inverted abstracts are reconstructed from indexed positions. The system translates filters separately because OpenAlex relevance search and PubMed Boolean syntax do not behave identically. Search responses retain the query used, sources searched, total counts, and structured paper metadata.

Citation snowballing uses a verified OpenAlex work resolved from a DOI, PMID, OpenAlex identifier, or exact normalized title. The result model marks whether a paper entered through database search, snowballing, or upload. Deduplication uses normalized DOI first and normalized title second; duplicate records are merged so that, for example, an abstract available from one source can enrich the retained record. The implementation records aggregate duplicates removed and PRISMA-oriented identification counts, but it does not preserve a complete user-visible lineage of every merged duplicate record. That remains an extension for stronger provenance auditing.

5.7 Protocol-Based Screening
Stage 5 requires an explicit set of inclusion and exclusion criteria. Default exclusion categories prevent an AI or human exclusion from being recorded without an available reason structure. Papers can be included, excluded, or marked uncertain/pending. The result record stores the current status, exclusion criterion, free-text screening reason, and whether the current decision was made by AI or a human.

The AI classifier is instructed to use the supplied criteria and can fetch a real abstract when the local record is insufficient. An uncertain AI result is routed into the review list as pending rather than silently excluded. Users can override AI decisions, clear AI-made screening decisions while leaving human decisions intact, add notes, and move papers among states. This actor distinction is one of the clearest operationalizations of human contribution: the final inclusion set can be inspected together with who made each current decision and why.

5.8 Literature and System Matrices
Stage 6 creates one literature row per retained paper and, optionally, one system row per selected tool or platform. Columns are user-controlled. AI autofill runs concurrently across rows and returns a value, evidence type, extraction method, confidence score, and confidence level. Each cell also includes an AI-origin flag, source-reference field, and human-verification field in the typed state model.

The current interface displays cell content and provides an on-demand evidence-and-reasoning popover. Opening that view is journaled as an AI-explanation/evidence interaction. However, the present autofill path initializes the source-reference field as empty, and the popover's explanation and sources are generated on demand rather than stored as immutable source spans. The implementation therefore supports inspection and records that inspection occurred, but it should not yet be described as complete passage-level provenance or as a fully implemented verification-checkbox workflow. Section 10 treats this as a priority limitation.

Users may also upload PDF or text files or provide a URL/DOI. The backend extracts PDF text, resolves OpenAlex metadata, attempts open-access PDF retrieval, and uses Europe PMC full-text XML as a biomedical fallback. Failures fall back to available abstract evidence rather than being presented as equivalent full-text extraction.

5.9 Evidence Synthesis
Stage 7 generates themes from the literature and optional system matrices. Each theme stores coverage, paper count, linked literature columns, linked system columns, and blocks for focus, patterns, unclear findings, missing evidence, or a user-created custom analysis. Users can edit block content and change which columns inform a theme, with selected changes recorded in the journal.

This design keeps the matrix as the central comparison surface. It does not prove that a generated synthesis claim is supported by every relevant cell. Rather, it narrows the input to structured columns and carries the chosen evidence dimensions forward, making the basis of a theme more inspectable than an unconstrained request over an undifferentiated document set.

5.10 Gap Identification
Stage 8 represents gaps as dense, thin, ignored, weak-design, uncertain, or no-theory dimensions and supports gap analysis for individual literature or system columns. The user selects candidate gaps and edits an overall gap statement. Journal entries can preserve the earlier and revised statement together with a justification.

The workflow explicitly separates a research gap from a missing cell. Sparse or absent evidence may indicate a genuine knowledge gap, inadequate searching, unsuitable schema design, or inaccessible full text. The user remains responsible for deciding which interpretation is defensible and whether additional retrieval is necessary.

5.11 Final Topic, Research Question, Study Design, and Structure
Stage 9 generates final topic candidates and alternative research questions from the selected gap and synthesis state. It retrieves a maintained catalog of study designs, uses the model to explain matches, and constructs a topic-specific paper structure from reporting-guideline anchor sections. Users choose the final topic, question, and design and can edit or reorder generated sections.

This stage demonstrates why upstream provenance matters. A polished question is not independently valid; it inherits the scope, search, selection, and extraction decisions made earlier. Matrica retains those upstream objects in the same project snapshot so that a reviewer can inspect the route to the final choice.

5.12 Cross-Cutting Human-Control and Contribution Model
Figure 2 summarizes how Matrica separates AI proposal, human intervention, recorded state, and evaluation. AI-origin markers identify generated suggestions. Screening records distinguish AI and human actors. Assistant-proposed changes are restricted to an opt-in catalog of typed actions, validated against current entity identifiers, and displayed as pending cards that the user can Apply or Skip; their outcomes remain applied, skipped, or failed. Selected user actions create journal entries with stage, entity, previous/new values, details, optional justification, and time. Evidence-view events show where the user requested support. A token-based reviewer workspace exposes each stage's results and journaled operations for comment without allowing the reviewer to alter the owner's decisions.

[[FIGURE_2]]
Figure 2. Contribution observability in Matrica. The record supports reconstruction and evaluation of human involvement but does not convert activity counts into a measure of intellectual quality.

Undo/redo retains up to 30 client states during a session. Autosave writes through to local browser storage and to a durable SQLite-backed project store, with local storage providing an offline fallback. A confirmed return to an earlier workflow stage marks later stages stale without deleting their content. This is a useful step-level warning, but it is not the fine-grained artifact dependency graph described in the earlier design specification. Current staleness is based on workflow position and back-navigation; automatic propagation from every upstream field edit to precisely affected downstream artifacts remains future work.

Taken together, these mechanisms make human contribution more visible than a final report alone. They still require interpretation. A long journal can document activity without demonstrating good reasoning, while a concise intervention may be decisive. The proposed evaluation therefore assesses the appropriateness and consequences of human actions, not merely their frequency.
"""


implementation = """6. System Architecture and Implementation
6.1 Runtime and Deployment Architecture
Matrica is implemented as a browser-based client and API service that can be deployed with Docker Compose. The frontend uses React 19, TypeScript 6, and Vite 8 and is served through Nginx. Nginx hosts the single-page application and proxies API requests to a FastAPI backend. The backend uses Pydantic request and response models, asynchronous HTTP calls for external scholarly services, and the OpenAI Responses API for model-assisted stages. The default model is configurable through an environment variable and falls back to GPT-4o in the inspected revision.

The deployment separates user-visible project data from the API key. Credentials remain in the server environment and are not included in exported project reports. The application can operate with local browser persistence when the server store is temporarily unreachable, while server-side project snapshots support recovery and reviewer sharing.

[[FIGURE_3]]
Figure 3. Implemented Matrica runtime architecture and data flows.

6.2 Frontend State and Interaction Architecture
Nine lazily loaded screen components are routed from a shared AppProvider. AppState is a typed representation of the entire project, including the Research Profile, schema, systems, topics, queries, search results, screening protocol, matrices, themes, gaps, final options, decision journal, and persistence metadata. A reducer applies discrete actions such as updating a profile field, excluding a paper, adding a matrix column, selecting a research question, or restoring a prior state.

The reducer architecture provides two safeguards. First, updates are expressed as named actions rather than arbitrary component mutation, which makes them testable and eligible for audit entries. Second, AI-operable actions are opt-in. Only action types tagged for chatbot use are included in a generated catalog. Before dispatch, the client checks required payload fields and rejects identifiers that do not exist in the current state. Potentially broad operations such as clearing all results, creating a new session, or restoring arbitrary state are not exposed to the assistant catalog.

The user interface adds a confirmation boundary. The assistant returns a conversational reply plus zero or more structured proposed actions. Each proposed action appears with Apply and Skip controls. The client stores pending, applied, skipped, or failed status. This ensures that conversational language alone does not silently change the research project and makes the user's decision to accept or reject an AI action visible.

6.3 Backend Research Services
The FastAPI backend is organized into routers for chat/profile formation, topic generation and reasoning, query generation, system recommendation, search, paper recognition, screening, matrix generation and ingestion, synthesis, gap analysis, final questions and study design, project storage, reviews, and analytics. Each stage validates structured inputs and parses model output into typed responses.

Retrieval uses OpenAlex and PubMed. Source-specific query preparation preserves differences between their search semantics, while filters are translated to each provider. PubMed XML is parsed to obtain real abstracts. OpenAlex supplies bibliographic metadata, abstracts where available, and citation networks. DOI- and title-based deduplication is implemented on both server and client paths. Paper ingestion accepts files and URLs/DOIs; PDF extraction uses pypdf, and open-access recovery checks OpenAlex locations and Europe PMC.

Screening accepts explicit inclusion and exclusion criteria and returns include, exclude, or uncertain plus a rationale and criterion identifier. Matrix services generate schemas, fill paper and system rows, retrieve on-demand evidence explanations, and accept document ingestion. Later services receive structured matrix and profile state rather than relying only on conversational history.

6.4 Persistence, Recovery, and Staleness
Client snapshots serialize the current project into a SessionSnapshot containing every workflow stage and the decision journal. Changes are autosaved after a short debounce to local storage and written through to the backend project endpoint. The server stores full JSON snapshots plus list metadata in SQLite. Existing local sessions are migrated to the server when possible.

Undo and redo use bounded in-memory state stacks. Session restore loads a complete snapshot. Back-navigation requires confirmation and marks every later workflow step stale while keeping existing content visible. Visiting a stale step clears its current stale marker. This implementation supports recovery and warns about broad downstream risk, but it does not calculate field-level dependency edges or prove that a stale result is incorrect.

6.5 Decision Journal and Human-Contribution Instrumentation
The decision journal records three classes of events: user operations, user decisions with optional justification, and views of AI evidence or explanation. Automatic entries cover profile changes, initial schema choices, system and topic selection, search filters, searches, screening decisions made by the user, synthesis evidence-column changes, and final selections. Stage components add more specific entries for query work, matrix schema changes, synthesis and gap edits, and evidence inspection.

Screening provenance is partly stored directly on each paper through screenedBy, screeningReason, exclusionReasonId, and uncertainty state. Matrix provenance is partly stored on each cell through AI origin, confidence, evidence type, extraction method, source reference, and verification fields. Assistant-action status provides a third record of human agency. These structures make it possible to reconstruct multiple kinds of intervention, although their coverage is not yet exhaustive or normalized into one event schema.

6.6 Peer Review, Analytics, and Export
The peer-review subsystem creates revocable bearer-token links. The reviewer receives a redacted project snapshot organized by all nine stages and can comment on each stage's result and on recorded user/AI-evidence operations. Reviewers cannot include, exclude, approve, or replace the user's decisions. Comments autosave, can be submitted, and are displayed to the owner with invitation status.

Workflow analytics use SQLite in write-ahead-log mode and per-thread connections. Middleware records instrumented API calls, durations, status codes, and inferred stage when a session identifier is available. Client events add session start, step changes, time on step, and periodic visible-tab heartbeats. The analytics dashboard exposes funnel, timing, regeneration, and potentially stuck-session views. These facilities make the planned evaluation technically feasible, but existing operational logs are not analyzed as participant evidence in this paper.

Export supports Word, browser-generated PDF, BibTeX, RIS, literature-matrix CSV, and PRISMA-flow CSV. The report export carries the Research Profile, topics, queries, retained papers, matrix, synthesis, gaps, questions, study design, and final structure. PRISMA-oriented CSV output includes identification counts, duplicate removal, screening totals, exclusion-reason breakdowns, and eligibility criteria.

6.7 Implementation Verification
Table 3 summarizes the inspected implementation and verification boundary. All automated checks were executed against the repository revision identified in Section 4.3.

Table 3. Implemented components and technical verification
Component	Implemented mechanism	Verification evidence
Client workflow	Nine React/TypeScript stages, typed reducer state, assistant action confirmation, decision journal, undo/redo, exports, reviewer and analytics views	42 frontend tests passed across five files; TypeScript project check completed; production bundle built successfully.
API and research services	FastAPI routers for generation, OpenAlex/PubMed search, citation chasing, screening, matrices, synthesis, gaps, final design, projects, reviews, and analytics	77 backend tests passed, including search, screening, review, agreement, and database-concurrency coverage.
Persistence and review	Local and server snapshots, SQLite project store, revocable review links, reviewer comments, WAL-mode analytics	Verified by source inspection and included backend/frontend test coverage; not evaluated with external users in this paper.
Evidence and contribution controls	Actor-marked screening, cell metadata, evidence-view journaling, optional justifications, AI action Apply/Skip status, step-level staleness	Mechanisms present in typed state and interfaces; passage-level immutable provenance and fine-grained dependency propagation remain incomplete.

Technical success means that the available tests pass and the application builds. It should not be interpreted as evidence that the workflow is easier to use, that AI outputs are accurate, or that recorded human activity reflects high-quality judgment. Those questions require the evaluation in Section 7.
"""


evaluation_intro = """7. Evaluation Protocol
7.1 Evaluation Objectives and Instrumentation
The evaluation is designed to complete the remaining DSRM evaluation activity and to test four claim classes that should not be conflated: implementation integrity, task-output reliability, effects on human effort, and the quality of human–AI collaboration. Implementation integrity is partly addressed by Section 6.7. The remaining claims require adjudicated benchmark data and participant studies.

The revised evaluation adds human-contribution outcomes to the earlier protocol. It asks not only whether a participant corrects an error, but also whether the participant inspects relevant evidence, supplies a defensible reason, preserves uncertainty when evidence is insufficient, recognizes when a gap may be caused by inadequate search, and responds appropriately to an upstream change. Journal and interface events provide observable traces, while interviews and expert review are required to interpret them. Event counts alone are not treated as evidence of critical thinking.

Three linked studies are proposed. A technical benchmark evaluates retrieval, screening, extraction, and source support against adjudicated reference data. A controlled crossover study compares Matrica with a manual structured-review workflow under equivalent evidence access. A mechanism study introduces seeded errors, ambiguous cases, and upstream changes to test error detection, appropriate reliance, recovery, and contribution observability. Existing analytics and journal structures support collection, but recruitment and analysis require ethics approval and informed consent.
"""


results_discussion = """8. Design and Implementation Results
8.1 Requirement-to-Mechanism Verification
The implemented prototype can be checked structurally against R1–R7. Table 5 distinguishes present mechanisms from claims that still require user or benchmark evidence.

Table 5. Requirement-to-mechanism implementation verification
Requirement	Implemented mechanism	Observable verification test
R1	Editable five-field Research Profile carried in shared context	Edit a field and confirm the new value and audit entry persist in the project snapshot.
R2	Typed SessionSnapshot spanning all nine stages	Save, reload, and inspect whether stage outputs reconstruct the project state.
R3	Queries, filters, criteria, status, reasons, uncertainty, and screenedBy	Trace an included, excluded, or uncertain paper to the current criterion, reason, and actor where present.
R4	User-controlled literature/system schemas and structured matrix cells	Compare rows under one schema and inspect stored evidence type, extraction method, and confidence metadata.
R5	Decision journal, assistant Apply/Skip status, evidence-view events, and reviewer comments	Recover selected human interventions and ask an independent reviewer to evaluate their methodological quality.
R6	Undo/redo, local/server snapshots, and step-level stale state	Return to an earlier stage, verify later steps are flagged, and restore a prior snapshot without deleting work.
R7	Structured transfer from matrices to synthesis, gaps, final topics, questions, design, and export	Trace a final choice to retained upstream project objects and exported report sections.

R1, R2, R3, R4, R6, and R7 have identifiable implementation surfaces. R5 is also implemented, but only for a selected set of interactions. Matrix evidence is inspectable, yet immutable source-span provenance and a visible end-user verification toggle are incomplete. Likewise, R6 is currently step-level rather than artifact-level. These qualifications prevent design intent from being reported as full implementation.

8.2 Technical Verification Outcome
The backend completed 77 automated tests successfully. The frontend completed 42 tests across five files. The TypeScript project check produced no errors, and the production client bundle built successfully. Together with source inspection, these results demonstrate that the repository contains an executable prototype with the architecture and major workflow mechanisms described in Section 6.

These checks are narrower than a benchmark. They do not measure the recall of a live search, the sensitivity of AI screening, the accuracy of an extracted matrix value, or the validity of a generated gap. External scholarly services and language models can also change independently of the repository. Technical verification should therefore be reported with the inspected revision, test date, configured model, and service conditions.

8.3 Human-Contribution Outcome at the Design Level
At the design level, Matrica makes several categories of human contribution observable. It records problem framing through profile edits; analytic framing through schema decisions; evidence-boundary choices through systems, topics, queries, filters, criteria, and paper decisions; interpretive choices through synthesis columns and gap revisions; and final design choices through selected questions, topics, study designs, and structures. It also records selected evaluative behaviors, including viewing AI evidence, applying or skipping assistant actions, and adding justifications.

This is a representational result, not a finding that users reason better. The journal can show that a person changed a criterion, but only the content and consequences of the change reveal whether it improved the review. Peer comments and mechanism-study outcomes are therefore necessary to evaluate contribution quality. Matrica's contribution is to make such evaluation possible within one project record.

8.4 Empirical Evidence Status and Falsification Conditions
No adjudicated benchmark, participant dataset, or completed comparative study accompanies the implementation. Retrieval coverage, screening sensitivity, extraction accuracy, completion time, SUS, NASA-TLX, appropriate-reliance scores, and improvement in scholarly quality are not reported. Existing analytics data are outside the present study and cannot substitute for a consented evaluation dataset.

The artifact would fail its intended contribution if users cannot reconstruct why a consequential artifact was produced; if AI and human decisions cannot be distinguished where the distinction matters; if evidence cues encourage acceptance without inspection; if step-level staleness is too coarse to support recovery; or if the journal produces activity traces that reviewers cannot interpret. The empirical claim would also fail if the integrated workflow increases burden without improving error detection, reporting, or recoverability. These conditions make the design revisable rather than self-validating.

9. Discussion
9.1 Answers to the Research Questions
RQ1 is answered by seven requirements that respond to the normalization and fallibility of AI assistance. Structured intent and explicit process state preserve the review context. Transparent selection and comparable evidence expose consequential decisions. Observable human agency makes intervention available for evaluation. Recoverable change protects against silent inconsistency, and evidence-aligned outcomes constrain downstream generation. The requirements shift the design target from an AI that produces an answer to a workflow that supports responsible use.

RQ2 is answered by the implemented nine-stage artifact. Matrica combines shared typed state, stage-specific services, actor-marked screening, matrix metadata, assistant action confirmation, journaled user operations, snapshots, peer review, analytics, and export. The implementation also reveals the limits of the earlier specification: staleness is currently step-level, and evidence explanations are not yet immutable source spans. Reporting these boundaries is itself part of provenance-aware system research.

RQ3 is answered by separating implementation integrity, output reliability, effort, appropriate reliance, and contribution quality. Automated tests show that mechanisms exist; benchmark metrics test whether outputs are correct; user measures test cost and usability; seeded errors test whether participants rely appropriately; and expert review of journaled interventions tests whether human involvement is methodologically meaningful. No single aggregate score can stand in for all five.

9.2 From Avoiding AI to Governing Its Use
The revised motivation begins from a practical observation: AI is increasingly present in literature-review work and cannot be treated only as an optional novelty. This does not justify mandatory adoption. It instead raises the standard for disclosure, evaluation, and interface design. When use cannot be assumed away, researchers need to know how to use AI within defensible boundaries, how to evaluate the result, and how to show what remains human work.

This shift matters because prohibition and automation are not the only choices. A workflow can allocate high-volume classification or drafting to a model while retaining human authority over scope, criteria, uncertainty, interpretation, and final claims. Matrica encodes this allocation in state and interaction records. The goal is not to prove that a human touched every output; it is to make the allocation visible enough to question.

9.3 Human Contribution Is More Than Approval
Human-in-the-loop systems can overstate control when the loop consists of repeated confirmation. Matrica attempts to broaden the visible human role by capturing framing, schema, selection, evidence inspection, justification, correction, and final choice. The peer reviewer is instructed to assess the user's critical-thinking process and cannot silently substitute independent decisions.

Nevertheless, observability is not equivalent to quality. More interactions may indicate diligence, confusion, poor interface design, or unnecessary friction. A researcher who correctly rejects one consequential false exclusion may contribute more than another who approves hundreds of routine cells. Future evaluation should therefore combine event traces with task context, expert judgment, and interviews. Contribution should be interpreted through the decisions made and their evidential consequences.

9.4 Reliability, Control Cost, and Appropriate Reliance
Explicit artifacts, prompts for reasons, confirmation cards, and stale warnings create interaction cost. That cost is justified only if it reduces consequential errors, improves recovery, makes reporting easier, or supports better peer scrutiny. The proposed crossover study measures time, actions, edits, and workload alongside error detection and output quality so that added friction cannot be mislabeled as control.

Uncertainty and provenance cues must be actionable. An uncertain screening decision should route a record to human judgment or better evidence. An evidence view should allow comparison with a source. A skipped assistant action should leave the project unchanged. A stale step should reveal why its validity is in question. The current implementation satisfies these principles unevenly: screening and action confirmation are relatively explicit, while matrix provenance and stale dependencies need further engineering.

9.5 Implications for Research AI
Four implications follow. First, Research AI should represent the researcher's epistemic work, not only model output. Second, provenance should be layered: source, transformation, and contribution provenance answer different questions. Third, evaluation should test the workflow in which outputs are used, because a model's task accuracy does not reveal whether users notice its failures. Fourth, implementation papers should distinguish an intended mechanism from a verified feature and a verified feature from an empirically beneficial outcome.

For reporting, researchers should disclose evidence sources, search dates and filters, model and product versions, relevant prompts or decision rules, human adjudication, missingness, and the nature of human corrections. Matrica can capture part of this information, but a final review remains responsible for presenting it in accordance with the applicable reporting and methodological guidance.

10. Limitations and Ethical Considerations
The first limitation is the absence of summative empirical evaluation. The repository and its automated tests were inspected, but no adjudicated benchmark or participant study was available. Passing tests demonstrate implementation integrity within their coverage; they do not demonstrate review validity, usability, or improved human reasoning.

Second, provenance coverage is incomplete. Matrix cells contain structured metadata, but the current autofill path does not persist immutable source passages, and the evidence popover generates reasoning and source suggestions on demand. A future implementation should store source identifiers, exact spans or page locations, retrieval time, transformation method, model and prompt version, and human verification event together. Until then, a generated explanation must not be treated as equivalent to primary evidence.

Third, contribution logging is selective. The decision journal covers important edits, selections, screening decisions, justifications, and evidence views, but not every consequential action is represented uniformly. Some provenance lives in entity fields, some in journal entries, and some in assistant-action status. A normalized event model would improve completeness and analysis. Conversely, more logging raises privacy and surveillance concerns and should be limited to methodologically justified events with user disclosure and consent.

Fourth, stale-state propagation is coarse. Confirmed back-navigation marks later steps stale, but ordinary upstream edits do not automatically traverse a fine-grained dependency graph. The system may therefore leave a downstream artifact appearing current after a specific upstream field changes unless the user navigates through the warning path. Artifact-level dependency edges and automatic invalidation should precede strong claims of dependency-aware provenance.

Fifth, retrieval is centered on OpenAlex and PubMed. Their disciplinary coverage, metadata, and query semantics differ, and neither substitutes for a review-specific multi-database search. Citation snowballing and uploads can extend the corpus, but coverage still depends on identifiers and accessible metadata. Open-access full text is incomplete, and PDF extraction can lose tables, images, equations, and layout-dependent meaning.

Sixth, model output depends on provider, model version, prompt, temperature, context, and safety behavior. Confidence fields are not calibrated probabilities. Refusals and missing output may be non-random [19]. Commercial products and external APIs also change, so evaluations must record dates, versions, tiers, and failures rather than treating runs months apart as equivalent.

Seventh, schema-first matrices can improve comparability while narrowing attention. Researchers may privilege what fits the columns and overlook emergent concepts. Synthesis and gap generation inherit search omissions, screening errors, and extraction weaknesses. Human verification reduces but does not eliminate these risks, and a person may reinforce rather than correct a model's bias.

Ethically, the workflow may contain unpublished research plans, participant or clinical information, copyrighted full text, and API credentials. Users require clear disclosure of what is transmitted to external services, how long it is retained, and what appears in reviewer links or exports. Review tokens are bearer credentials and must be revocable, protected, and scoped. Logs should support accountability without becoming an unnecessary repository of sensitive behavioral data.

Finally, automation changes scholarly learning and responsibility. Screening, coding, and discussion can build domain knowledge and reviewer consensus [10]. If AI prefilters those activities, evaluation should examine learning, critical engagement, and accountability in addition to speed. Human control is meaningful only when users have access to evidence, authority to disagree, and enough time to exercise judgment.

11. Conclusion
AI assistance is becoming difficult to avoid in many literature-review settings, but inevitability is not a justification for uncritical use. It is a reason to make use governable. The important questions are how the tool is used, what evidence it sees, what the researcher decides, how uncertain or incorrect output is handled, and whether the resulting claim can be reconstructed and challenged.

Matrica addresses these questions through an implemented nine-stage workflow that carries structured state from a Research Profile to search, screening, matrices, synthesis, gaps, and a final research design. Its distinctive emphasis is the visibility of collaboration: actor-marked screening, user-controlled schemas, reasons and edits, evidence-view records, human-confirmed assistant actions, project snapshots, peer process comments, and exports that retain intermediate artifacts.

The implementation is real and technically verified within the available checks: 77 backend tests and 42 frontend tests pass, TypeScript checking succeeds, and the production client builds. The empirical value of the artifact is not yet established. Matrica has not been shown in this paper to improve accuracy, efficiency, usability, trust, or critical thinking. Matrix provenance and dependency handling also remain less complete than the original design specification.

The next research step is therefore evaluation, not initial implementation. Adjudicated benchmark reviews and a preregistered user study should test whether Matrica helps researchers detect and repair consequential errors, preserve uncertainty, explain their decisions, and recover from upstream change without unacceptable overhead. If it does, Matrica will contribute a more accountable model of human-controlled Research AI. If it does not, the explicit state, interaction records, and falsification conditions provide a basis for revision.

Software and Data Availability
The Matrica source code is available in the TPP Projects GitHub repository [45]. The implementation claims in this paper refer to revision 9e30bc4. No benchmark or participant dataset is associated with the present manuscript.
"""


manuscript = "\n\n".join(
    [
        front_matter.strip(),
        "3. Background and Related Work\n" + related_work_31_33,
        human_ai_related.strip(),
        methodology.strip(),
        framework.strip(),
        implementation.strip(),
        evaluation_intro.strip(),
        evaluation_72_76,
        results_discussion.strip(),
        references,
    ]
)


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(FONT_BOLD if bold else FONT_REGULAR, size=size)


def rounded_box(draw, xy, fill, outline=BLUE, radius=20, width=3):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)


def fit_lines(draw, text: str, fnt, max_width: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    line = ""
    for word in words:
        trial = f"{line} {word}".strip()
        if draw.textlength(trial, font=fnt) <= max_width or not line:
            line = trial
        else:
            lines.append(line)
            line = word
    if line:
        lines.append(line)
    return lines


def centered_text(draw, xy, text, fnt, fill=BLACK, max_width=None, spacing=8):
    x1, y1, x2, y2 = xy
    width = x2 - x1
    lines = fit_lines(draw, text, fnt, max_width or width - 30)
    bbox = fnt.getbbox("Ag")
    line_h = bbox[3] - bbox[1]
    total_h = len(lines) * line_h + max(0, len(lines) - 1) * spacing
    y = y1 + (y2 - y1 - total_h) / 2
    for line in lines:
        tw = draw.textlength(line, font=fnt)
        draw.text((x1 + (width - tw) / 2, y), line, font=fnt, fill=fill)
        y += line_h + spacing


def arrow(draw, x1, y1, x2, y2, color=BLUE, width=5):
    draw.line((x1, y1, x2, y2), fill=color, width=width)
    if abs(x2 - x1) >= abs(y2 - y1):
        direction = 1 if x2 > x1 else -1
        pts = [(x2, y2), (x2 - direction * 16, y2 - 10), (x2 - direction * 16, y2 + 10)]
    else:
        direction = 1 if y2 > y1 else -1
        pts = [(x2, y2), (x2 - 10, y2 - direction * 16), (x2 + 10, y2 - direction * 16)]
    draw.polygon(pts, fill=color)


def draw_figure_1(path: Path):
    img = Image.new("RGB", (1800, 820), "white")
    d = ImageDraw.Draw(img)
    title_f = font(36, True)
    box_f = font(25, True)
    small_f = font(21)
    d.text((70, 35), "From research intent to an inspectable research design", font=title_f, fill=DARK_BLUE)
    labels = [
        ("1", "Define topic"), ("2", "Matrix initial"), ("3", "System review"),
        ("4", "Explore & search"), ("5", "Select papers"),
        ("6", "Matrix table"), ("7", "Synthesis"), ("8", "Gap & statement"),
        ("9", "Final question"),
    ]
    positions = []
    for i, item in enumerate(labels):
        row = 0 if i < 5 else 1
        col = i if row == 0 else i - 5
        x = 70 + col * 338
        y = 135 + row * 250
        w = 280
        h = 130
        positions.append((x, y, x + w, y + h))
        rounded_box(d, positions[-1], LIGHT_BLUE if i % 2 == 0 else LIGHT_GRAY)
        d.ellipse((x + 15, y + 15, x + 61, y + 61), fill=BLUE)
        centered_text(d, (x + 15, y + 15, x + 61, y + 61), item[0], font(23, True), fill=WHITE)
        centered_text(d, (x + 65, y + 12, x + w - 10, y + h - 12), item[1], box_f, fill=DARK_BLUE)
    for i in range(4):
        arrow(d, positions[i][2] + 8, (positions[i][1] + positions[i][3]) // 2, positions[i + 1][0] - 8, (positions[i + 1][1] + positions[i + 1][3]) // 2)
    arrow(d, positions[4][2] - 45, positions[4][3] + 8, positions[5][2] - 45, positions[5][1] - 8)
    for i in range(5, 8):
        arrow(d, positions[i][2] + 8, (positions[i][1] + positions[i][3]) // 2, positions[i + 1][0] - 8, (positions[i + 1][1] + positions[i + 1][3]) // 2)
    band = (70, 665, 1730, 770)
    rounded_box(d, band, "#FFF7E8", outline=GOLD, radius=18, width=3)
    centered_text(
        d, band,
        "Human contribution across stages: define • select • inspect • edit • justify • confirm • review",
        small_f, fill="#6B4A10",
    )
    img.save(path, dpi=(220, 220))


def draw_figure_2(path: Path):
    img = Image.new("RGB", (1800, 760), "white")
    d = ImageDraw.Draw(img)
    title_f = font(36, True)
    head_f = font(27, True)
    text_f = font(22)
    d.text((70, 35), "How human–AI collaboration becomes observable", font=title_f, fill=DARK_BLUE)
    boxes = [
        (70, 160, 390, 360, "AI proposes", "topic • query • classification • cell • synthesis • gap"),
        (520, 160, 840, 360, "Human intervenes", "select • inspect • edit • justify • apply/skip • override"),
        (970, 160, 1290, 360, "System records", "state • actor • reason • prior/new value • time • status"),
        (1410, 160, 1730, 360, "Reviewer evaluates", "result quality • critical thinking • appropriateness"),
    ]
    fills = ["#EEF3FA", "#FFF7E8", "#EEF7F2", "#F4F0FA"]
    outlines = [BLUE, GOLD, "#397357", "#6B4E8E"]
    for (x1, y1, x2, y2, head, body), fill, outline in zip(boxes, fills, outlines):
        rounded_box(d, (x1, y1, x2, y2), fill, outline=outline)
        centered_text(d, (x1 + 15, y1 + 12, x2 - 15, y1 + 75), head, head_f, fill=outline)
        centered_text(d, (x1 + 20, y1 + 78, x2 - 20, y2 - 15), body, text_f, fill=BLACK)
    for i in range(3):
        arrow(d, boxes[i][2] + 10, 260, boxes[i + 1][0] - 10, 260)
    rounded_box(d, (240, 505, 1560, 660), LIGHT_GRAY, outline=MID_GRAY, radius=18, width=2)
    centered_text(
        d, (260, 520, 1540, 645),
        "Interpretation remains necessary: an event trace documents participation, but methodological quality depends on the decision, evidence, and consequence.",
        font(24), fill=BLACK,
    )
    img.save(path, dpi=(220, 220))


def draw_figure_3(path: Path):
    img = Image.new("RGB", (1800, 900), "white")
    d = ImageDraw.Draw(img)
    d.text((70, 35), "Implemented runtime architecture", font=font(36, True), fill=DARK_BLUE)
    layers = [
        (120, 135, 1680, 265, "Browser client", "React + TypeScript • nine workflow screens • typed reducer • confirmation cards • exports"),
        (120, 325, 1680, 455, "API and orchestration", "FastAPI • stage routers • structured request/response models • OpenAI Responses API"),
        (120, 515, 1680, 645, "Evidence and document services", "OpenAlex • PubMed • citation snowballing • PDF/URL ingestion • Europe PMC fallback"),
        (120, 705, 1680, 835, "State, review, and observability", "local snapshots • SQLite projects/analytics • peer-review tokens • decision journal"),
    ]
    fills = ["#EEF3FA", "#F3F0FA", "#EEF7F2", "#FFF7E8"]
    outlines = [BLUE, "#6B4E8E", "#397357", GOLD]
    for (x1, y1, x2, y2, head, body), fill, outline in zip(layers, fills, outlines):
        rounded_box(d, (x1, y1, x2, y2), fill, outline=outline)
        d.text((x1 + 35, y1 + 25), head, font=font(28, True), fill=outline)
        d.text((x1 + 35, y1 + 72), body, font=font(22), fill=BLACK)
    for i in range(3):
        arrow(d, 900, layers[i][3] + 8, 900, layers[i + 1][1] - 8)
    img.save(path, dpi=(220, 220))


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa: int):
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn("w:tcW"))
    if tcW is None:
        tcW = OxmlElement("w:tcW")
        tcPr.append(tcW)
    tcW.set(qn("w:w"), str(width_dxa))
    tcW.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]):
    total = sum(widths)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(total))
    tblW.set(qn("w:type"), "dxa")
    tblInd = tblPr.find(qn("w:tblInd"))
    if tblInd is None:
        tblInd = OxmlElement("w:tblInd")
        tblPr.append(tblInd)
    tblInd.set(qn("w:w"), "120")
    tblInd.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths[min(idx, len(widths) - 1)]
            set_cell_width(cell, width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_shading(cell, fill: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill.lstrip("#"))


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.rFonts
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color.lstrip("#"))
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, size=9, color=MID_GRAY)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def keep_with_next(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    if pPr.find(qn("w:keepNext")) is None:
        pPr.append(OxmlElement("w:keepNext"))


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


def widths_for_table(headers: list[str]) -> list[int]:
    n = len(headers)
    if n == 2:
        return [2700, 6660]
    if n == 3:
        first = headers[0].lower()
        if first in {"id", "requirement"}:
            return [900, 2350, 6110]
        if first == "component":
            return [1900, 3000, 4460]
        if first == "stage":
            return [1500, 3600, 4260]
        return [1800, 3300, 4260]
    return [9360 // n] * (n - 1) + [9360 - (9360 // n) * (n - 1)]


def add_table(doc: Document, caption: str, rows: list[list[str]]):
    cap = doc.add_paragraph(style="Caption")
    cap.add_run(caption)
    keep_with_next(cap)
    cols = max(len(r) for r in rows)
    normalized = [r + [""] * (cols - len(r)) for r in rows]
    table = doc.add_table(rows=len(normalized), cols=cols)
    table.style = "Table Grid"
    widths = widths_for_table(normalized[0])
    set_table_geometry(table, widths)
    for r_idx, data in enumerate(normalized):
        row = table.rows[r_idx]
        if r_idx == 0:
            set_repeat_table_header(row)
        for c_idx, value in enumerate(data):
            cell = row.cells[c_idx]
            if r_idx == 0:
                set_cell_shading(cell, BLUE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.05
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(value)
            set_run_font(run, size=8.5, color=WHITE if r_idx == 0 else BLACK, bold=(r_idx == 0))
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(4)


def add_body_paragraph(doc: Document, text: str):
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run(text)
    return p


def add_rq_paragraph(doc: Document, text: str):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.right_indent = Inches(0.2)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(5)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), LIGHT_BLUE.lstrip("#"))
    pPr.append(shd)
    prefix, rest = text.split(". ", 1)
    r1 = p.add_run(prefix + ". ")
    set_run_font(r1, bold=True, color=DARK_BLUE)
    p.add_run(rest)


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1.15)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.3)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BLACK.lstrip("#"))
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    pf = normal.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.space_before = Pt(0)
    pf.space_after = Pt(8)
    pf.line_spacing = 1.333
    pf.widow_control = True

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color.lstrip("#"))
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    caption = styles["Caption"]
    caption.font.name = "Calibri"
    caption.font.size = Pt(9)
    caption.font.bold = True
    caption.font.italic = False
    caption.font.color.rgb = RGBColor.from_string(DARK_BLUE.lstrip("#"))
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(4)
    caption.paragraph_format.keep_with_next = True

    if "Reference" not in [s.name for s in styles]:
        ref_style = styles.add_style("Reference", WD_STYLE_TYPE.PARAGRAPH)
    else:
        ref_style = styles["Reference"]
    ref_style.font.name = "Calibri"
    ref_style.font.size = Pt(9)
    ref_style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    ref_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    ref_style.paragraph_format.left_indent = Inches(0.3)
    ref_style.paragraph_format.first_line_indent = Inches(-0.3)
    ref_style.paragraph_format.space_after = Pt(4)
    ref_style.paragraph_format.line_spacing = 1.1

    header = section.header
    hp = header.paragraphs[0]
    hp.text = "MATRICA  |  Human–AI collaboration in literature reviews"
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in hp.runs:
        set_run_font(run, size=8.5, color=MID_GRAY, bold=True)
    footer = section.footer
    add_page_field(footer.paragraphs[0])


def add_title_page(doc: Document):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(78)
    p.paragraph_format.space_after = Pt(10)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("MATRICA")
    set_run_font(r, size=12, color=GOLD, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(16)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Matrica: A Provenance-Aware Workflow for Making Human Contribution Visible in AI-Assisted Literature Reviews")
    set_run_font(r, size=26, color=DARK_BLUE, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(40)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Implemented system paper and evaluation protocol")
    set_run_font(r, size=13, color=MID_GRAY, italic=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(100)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("Revision date: 3 August 2026")
    set_run_font(r, size=10, color=MID_GRAY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Implementation revision inspected: 9e30bc4")
    set_run_font(r, size=10, color=MID_GRAY)

    doc.add_page_break()


def render_manuscript(doc: Document, text: str, figures: dict[str, Path]):
    figure_alt_text = {
        "[[FIGURE_1]]": (
            "Nine-stage Matrica workflow from defining a topic through matrix construction, "
            "synthesis, gap analysis, and final question selection, with human contribution "
            "activities spanning every stage."
        ),
        "[[FIGURE_2]]": (
            "Human-AI collaboration model in which AI proposals are followed by human "
            "intervention, system recording, and reviewer evaluation."
        ),
        "[[FIGURE_3]]": (
            "Matrica runtime architecture linking the browser client, FastAPI orchestration, "
            "evidence and document services, and state, review, and observability services."
        ),
    }
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line in figures:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(3)
            picture = p.add_run().add_picture(str(figures[line]), width=Inches(6.45))
            picture._inline.docPr.set("descr", figure_alt_text[line])
            picture._inline.docPr.set("title", f"Matrica {line[2:-2].lower().replace('_', ' ')}")
            i += 1
            continue
        if re.match(r"^Table \d+\.", line):
            rows: list[list[str]] = []
            i += 1
            while i < len(lines):
                candidate = lines[i].strip()
                if not candidate:
                    i += 1
                    if rows:
                        break
                    continue
                if "\t" not in lines[i]:
                    break
                rows.append([part.strip() for part in lines[i].split("\t")])
                i += 1
            add_table(doc, line, rows)
            continue
        if re.match(r"^Figure \d+\.", line):
            p = doc.add_paragraph(style="Caption")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run(line)
            i += 1
            continue
        if line == "References":
            doc.add_paragraph(line, style="Heading 1")
            i += 1
            continue
        if line == "Software and Data Availability":
            doc.add_paragraph(line, style="Heading 1")
            i += 1
            continue
        if re.match(r"^\d+\.\d+\s", line):
            doc.add_paragraph(line, style="Heading 2")
            i += 1
            continue
        if re.match(r"^\d+\.\s", line):
            doc.add_paragraph(line, style="Heading 1")
            i += 1
            continue
        if re.match(r"^RQ\d+\.\s", line):
            add_rq_paragraph(doc, line)
            i += 1
            continue
        if line.startswith("Keywords:"):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after = Pt(10)
            first, rest = line.split(":", 1)
            r = p.add_run(first + ":")
            set_run_font(r, bold=True, color=DARK_BLUE)
            p.add_run(rest)
            i += 1
            continue
        if re.match(r"^\[\d+\]", line):
            p = doc.add_paragraph(style="Reference")
            p.add_run(line)
            i += 1
            continue
        add_body_paragraph(doc, line)
        i += 1


def main():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    fig1 = ASSET_DIR / "figure_1_workflow.png"
    fig2 = ASSET_DIR / "figure_2_contribution.png"
    fig3 = ASSET_DIR / "figure_3_architecture.png"
    draw_figure_1(fig1)
    draw_figure_2(fig2)
    draw_figure_3(fig3)

    doc = Document()
    configure_document(doc)
    add_title_page(doc)
    render_manuscript(
        doc,
        manuscript,
        {"[[FIGURE_1]]": fig1, "[[FIGURE_2]]": fig2, "[[FIGURE_3]]": fig3},
    )
    doc.core_properties.title = "Matrica: Making Human Contribution Visible in AI-Assisted Literature Reviews"
    doc.core_properties.subject = "Provenance-aware human–AI collaboration in structured literature review"
    doc.core_properties.author = "TPP Projects"
    doc.core_properties.keywords = "Research AI; systematic review; human–AI collaboration; provenance; human contribution"
    doc.core_properties.comments = "Revised against implementation revision 9e30bc4."
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
